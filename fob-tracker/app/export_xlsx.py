"""Monthly Excel export using openpyxl."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import report


def _header(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1E8A8A")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _hhmm(hours: float) -> str:
    total_min = int(round(hours * 60))
    return f"{total_min // 60:02d}:{total_min % 60:02d}"


def build_monthly_workbook(year: int, month: int) -> bytes:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_headers = [
        "Name",
        "Employee ID",
        "Days Worked",
        "Total Hours",
        "Total Hours (HH:MM)",
        "# Manual Entries",
    ]
    for col, name in enumerate(summary_headers, start=1):
        c = summary_ws.cell(row=1, column=col, value=name)
        _header(c)

    for idx, row in enumerate(report.month_summary(year, month), start=2):
        summary_ws.cell(row=idx, column=1, value=row["name"])
        summary_ws.cell(row=idx, column=2, value=row["employee_id"] or "")
        summary_ws.cell(row=idx, column=3, value=row["days_worked"])
        summary_ws.cell(row=idx, column=4, value=row["total_hours"])
        summary_ws.cell(row=idx, column=5, value=_hhmm(row["total_hours"]))
        summary_ws.cell(row=idx, column=6, value=row["manual_entries"])

    for col in range(1, len(summary_headers) + 1):
        summary_ws.column_dimensions[get_column_letter(col)].width = 20

    daily_ws = wb.create_sheet("Daily")
    daily_headers = [
        "Date",
        "Name",
        "Clock In",
        "Clock Out",
        "Session Hours",
        "Source",
        "Notes",
    ]
    for col, name in enumerate(daily_headers, start=1):
        c = daily_ws.cell(row=1, column=col, value=name)
        _header(c)

    for idx, row in enumerate(report.daily_pairs(year, month), start=2):
        daily_ws.cell(row=idx, column=1, value=row["date"])
        daily_ws.cell(row=idx, column=2, value=row["name"])
        daily_ws.cell(row=idx, column=3, value=row["clock_in"] or "")
        daily_ws.cell(row=idx, column=4, value=row["clock_out"] or "")
        daily_ws.cell(row=idx, column=5, value=row["hours"])
        daily_ws.cell(row=idx, column=6, value=row["source"])
        daily_ws.cell(row=idx, column=7, value=row["notes"])

    for col in range(1, len(daily_headers) + 1):
        daily_ws.column_dimensions[get_column_letter(col)].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

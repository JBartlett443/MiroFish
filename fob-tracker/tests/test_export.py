from io import BytesIO

from openpyxl import load_workbook

from app import export_xlsx
from tests.conftest import seed_user, seed_scan


def test_workbook_shape(ctx):
    uid = seed_user("Jane", "E1")
    seed_scan(uid, "2026-10-01 09:00:00", "IN")
    seed_scan(uid, "2026-10-01 17:00:00", "OUT")
    seed_scan(uid, "2026-10-02 09:00:00", "IN", source="manual")
    seed_scan(uid, "2026-10-02 17:00:00", "OUT", source="manual")

    data = export_xlsx.build_monthly_workbook(2026, 10)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Summary", "Daily"]

    summary = wb["Summary"]
    assert [c.value for c in summary[1]] == [
        "Name", "Employee ID", "Days Worked",
        "Total Hours", "Total Hours (HH:MM)", "# Manual Entries",
    ]
    row = [c.value for c in summary[2]]
    assert row[0] == "Jane"
    assert row[1] == "E1"
    assert row[2] == 2
    assert row[3] == 16.0
    assert row[4] == "16:00"
    assert row[5] == 2

    daily = wb["Daily"]
    assert [c.value for c in daily[1]] == [
        "Date", "Name", "Clock In", "Clock Out",
        "Session Hours", "Source", "Notes",
    ]
    sources = {daily.cell(row=r, column=6).value for r in range(2, daily.max_row + 1)}
    assert "manual" in sources
    assert "fob" in sources
